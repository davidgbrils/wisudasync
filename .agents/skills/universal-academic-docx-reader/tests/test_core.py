from pathlib import Path
import json
import pytest
from docx import Document
from docx_reader.core import extract_docx, inspect_docx, search_docx, list_docx_chunks, read_docx_chunk


def make_docx(path: Path):
    doc = Document()
    doc.core_properties.title = "Academic Reader Test"
    doc.add_heading("BAB I PENDAHULUAN", level=1)
    doc.add_paragraph("Penelitian ini membahas sistem informasi akademik.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Komponen"
    table.cell(0, 1).text = "Hasil"
    table.cell(1, 0).text = "Pengujian"
    table.cell(1, 1).text = "Berhasil"
    doc.sections[0].header.paragraphs[0].text = "Header Uji"
    doc.sections[0].footer.paragraphs[0].text = "Footer Uji"
    doc.add_paragraph(("Isi tambahan untuk pengujian chunk. " * 80))
    doc.save(path)


def test_extract_and_search(tmp_path):
    path = tmp_path / "sample.docx"
    make_docx(path)
    result = extract_docx(path, output_dir=tmp_path / "out", chunk_chars=800)
    assert result["ok"]
    out = Path(result["output_dir"])
    sections = json.loads((out / "sections.json").read_text(encoding="utf-8"))
    assert "sistem informasi akademik" in sections["body"]
    assert "Header Uji" in sections["header"]
    assert "Footer Uji" in sections["footer"]
    assert len(json.loads((out / "chunks.json").read_text(encoding="utf-8"))) >= 1
    search = search_docx(path, "sistem informasi akademik", output_dir=tmp_path / "out")
    assert search["result_count"] >= 1


def test_chunk_read(tmp_path):
    path = tmp_path / "sample.docx"
    make_docx(path)
    extract_docx(path, output_dir=tmp_path / "out", chunk_chars=800)
    listing = list_docx_chunks(path, output_dir=tmp_path / "out")
    assert listing["total_chunks"] >= 1
    first = read_docx_chunk(path, listing["chunks"][0]["chunk_id"], output_dir=tmp_path / "out")
    assert first["text"]


def test_invalid_extension(tmp_path):
    path = tmp_path / "bad.txt"
    path.write_text("not docx")
    with pytest.raises(ValueError):
        inspect_docx(path)
