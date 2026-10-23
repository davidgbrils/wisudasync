from docx import Document
from pathlib import Path

out = Path(__file__).with_name("sample_academic.docx")
doc = Document()
doc.core_properties.title = "Contoh Laporan Akademik"
doc.add_heading("BAB I PENDAHULUAN", level=1)
doc.add_paragraph("Latar belakang penelitian menjelaskan masalah dan kebutuhan sistem.")
doc.add_heading("1.1 Rumusan Masalah", level=2)
doc.add_paragraph("Bagaimana membangun sistem informasi yang dapat diuji secara terukur?")
doc.add_heading("BAB II METODOLOGI", level=1)
doc.add_paragraph("Metode penelitian menggunakan observasi, studi pustaka, dan pengujian.")
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Tahap"
table.cell(0, 1).text = "Keluaran"
table.cell(1, 0).text = "Pengujian"
table.cell(1, 1).text = "Laporan hasil uji"
doc.sections[0].header.paragraphs[0].text = "Contoh Header Akademik"
doc.sections[0].footer.paragraphs[0].text = "Contoh Footer Akademik"
doc.save(out)
print(f"Created: {out}")
