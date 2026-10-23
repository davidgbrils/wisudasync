from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook

from app.api.v1.endpoints import sitasi as sitasi_endpoint
from app.services import import_service
from app.services.import_service import (
    ImportValidationError,
    UploadTooLargeError,
    import_master_data,
    parse_certiport,
    parse_sitasi,
)


SITASI_HEADERS = [
    "NIM",
    "Nama",
    "Fakultas",
    "Kode NIM",
    "Program Studi",
    "Angkatan",
    "Sertifikasi",
    "Sub Program",
    "Nilai",
    "Status",
]

CERTIPORT_HEADERS = [
    "First Name",
    "Last Name",
    "Full Name",
    "Program Name",
    "Certification Name",
    "Score",
    "Result",
]


def make_xlsx(
    sheet_name: str,
    headers: list[str],
    rows: list[list[object]],
    *,
    header_row: int = 1,
    nim_number_format: str | None = None,
    formula: bool = False,
    formula_column: int | None = None,
    formula_value: str = "=1+1",
) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for _ in range(header_row - 1):
        sheet.append([])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if nim_number_format:
        sheet.cell(row=header_row + 1, column=1).number_format = nim_number_format
    if formula:
        sheet.cell(row=header_row + 1, column=2, value="=1+1")
    if formula_column is not None:
        for offset in range(len(rows)):
            sheet.cell(row=header_row + 1 + offset, column=formula_column, value=formula_value)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def sitasi_fixture(rows: list[list[object]], **kwargs) -> BytesIO:
    return make_xlsx("Master_Data", SITASI_HEADERS, rows, **kwargs)


def certiport_fixture(rows: list[list[object]], *, full_name_formula: bool = False, **kwargs) -> BytesIO:
    if full_name_formula:
        kwargs.setdefault("formula_column", 3)
        kwargs.setdefault("formula_value", '=A5&" "&B5')
    return make_xlsx(
        "Results (264) Plus",
        CERTIPORT_HEADERS,
        rows,
        header_row=4,
        **kwargs,
    )


def sitasi_row(nim: object, name: str, status: str) -> list[object]:
    return [nim, name, "Fakultas Uji", "KODE", "Prodi Uji", "2024", "MOS", "Word", 700, status]


class FakeQuery:
    def __init__(self, client: "FakeSupabase", table: str):
        self.client = client
        self.table = table
        self.operation = "select"
        self.payload = None
        self.filters: dict[str, object] = {}

    def insert(self, payload, **kwargs):
        self.operation = "insert"
        self.payload = payload
        return self

    def select(self, columns):
        self.operation = "select" if self.operation == "select" else self.operation
        return self

    def update(self, payload):
        self.operation = "update"
        self.payload = payload
        return self

    def delete(self):
        self.operation = "delete"
        return self

    def eq(self, column, value):
        self.filters[column] = value
        return self

    def order(self, *args, **kwargs):
        return self

    def limit(self, *args, **kwargs):
        return self

    def execute(self):
        if self.operation == "insert" and self.table == "data_imports":
            import_id = f"batch-{len(self.client.imports)}"
            self.client.imports.append({**self.payload, "id": import_id})
            return SimpleNamespace(data=[{"id": import_id}])

        if self.operation == "insert":
            self.client.row_insert_calls += 1
            if self.client.fail_row_insert_call == self.client.row_insert_calls:
                raise OSError("synthetic database detail with student data")
            self.client.rows.extend(self.payload)
            return SimpleNamespace(data=[])

        if self.operation == "delete":
            self.client.rows = [
                row
                for row in self.client.rows
                if not all(row.get(key) == value for key, value in self.filters.items())
            ]
            return SimpleNamespace(data=[])

        if self.operation == "update":
            for record in self.client.imports:
                if all(record.get(key) == value for key, value in self.filters.items()):
                    record.update(self.payload)
            return SimpleNamespace(data=[])

        records = [
            record
            for record in self.client.imports
            if all(record.get(key) == value for key, value in self.filters.items())
        ]
        return SimpleNamespace(data=records)


class FakeRpc:
    def __init__(self, client: "FakeSupabase", parameters: dict):
        self.client = client
        self.parameters = parameters

    def execute(self):
        if self.client.fail_rpc:
            raise OSError("synthetic RPC detail with student data")
        target_id = self.parameters["p_import_id"]
        for record in self.client.imports:
            if record["data_type"] == "SITASI":
                record["is_active"] = record["id"] == target_id
                if record["id"] == target_id:
                    record["status"] = "COMPLETED"
                    record["imported_rows"] = self.parameters["p_imported_rows"]
        return SimpleNamespace(data=[])


class FakeSupabase:
    def __init__(self):
        self.imports = [
            {"id": "old-batch", "data_type": "SITASI", "status": "COMPLETED", "is_active": True}
        ]
        self.rows: list[dict] = []
        self.row_insert_calls = 0
        self.fail_row_insert_call: int | None = None
        self.fail_rpc = False

    def table(self, table: str):
        return FakeQuery(self, table)

    def rpc(self, name: str, parameters: dict):
        assert name == "complete_data_import"
        return FakeRpc(self, parameters)


def test_sitasi_preserves_leading_zeroes_names_status_and_duplicate_attempts():
    workbook = sitasi_fixture(
        [
            sitasi_row("001234567", "  Nama   Uji ", "Lulus"),
            sitasi_row("001234567", "Nama Uji", "Gagal"),
        ]
    )

    parsed = parse_sitasi(workbook, "fixture.xlsx")

    assert [row["nim"] for row in parsed.rows] == ["001234567", "001234567"]
    assert [row["nama"] for row in parsed.rows] == ["  Nama   Uji ", "Nama Uji"]
    assert [row["nama_normalized"] for row in parsed.rows] == ["NAMA UJI", "NAMA UJI"]
    assert [row["status"] for row in parsed.rows] == ["LULUS", "GAGAL"]
    assert [row["source_row"] for row in parsed.rows] == [2, 3]


def test_sitasi_numeric_nim_uses_zero_mask_format():
    parsed = parse_sitasi(
        sitasi_fixture([sitasi_row(12345, "Nama Uji", "Lulus")], nim_number_format="00000000"),
        "fixture.xlsx",
    )

    assert parsed.rows[0]["nim"] == "00012345"


def test_certiport_finds_header_on_row_four_and_keeps_pass_fail():
    parsed = parse_certiport(
        certiport_fixture(
            [
                ["Ayu", "Uji", "Ayu Uji", "Microsoft Office Specialist", "Word 2019", 800, "Pass"],
                ["Bima", "Uji", "Bima Uji", "Microsoft Office Specialist", "Word 2019", 400, "Fail"],
            ]
        ),
        "fixture.xlsx",
    )

    assert [row["result"] for row in parsed.rows] == ["PASS", "FAIL"]
    assert [row["source_row"] for row in parsed.rows] == [5, 6]
    assert parsed.rows[0]["full_name_normalized"] == "AYU UJI"


@pytest.mark.parametrize(
    ("sheet", "headers", "message"),
    [
        ("Wrong_Sheet", SITASI_HEADERS, "Sheet wajib"),
        ("Master_Data", SITASI_HEADERS[:-1] + ["Result"], "header wajib"),
    ],
)
def test_sitasi_rejects_wrong_sheet_or_headers(sheet, headers, message):
    workbook = make_xlsx(sheet, headers, [sitasi_row("001", "Nama Uji", "Lulus")])

    with pytest.raises(ImportValidationError, match=message):
        parse_sitasi(workbook, "fixture.xlsx")


def test_import_rejects_formula_workbook():
    with pytest.raises(ImportValidationError, match="formula"):
        parse_sitasi(
            sitasi_fixture([sitasi_row("001", "Nama Uji", "Lulus")], formula=True),
            "fixture.xlsx",
        )


def test_certiport_accepts_full_name_formula_and_derives_from_first_last():
    parsed = parse_certiport(
        certiport_fixture(
            [["Ayu", "Uji", None, "Microsoft Office Specialist", "Word 2019", 800, "Pass"]],
            full_name_formula=True,
        ),
        "fixture.xlsx",
    )

    assert parsed.rows[0]["full_name"] == "Ayu Uji"
    assert parsed.rows[0]["full_name_normalized"] == "AYU UJI"


def test_certiport_rejects_formula_in_required_column():
    with pytest.raises(ImportValidationError, match="formula"):
        parse_certiport(
            certiport_fixture(
                [["Ayu", "Uji", "Ayu Uji", "Microsoft Office Specialist", "Word 2019", 800, "Pass"]],
                formula_column=6,
            ),
            "fixture.xlsx",
        )


def test_import_rejects_macro_member():
    source = sitasi_fixture([sitasi_row("001", "Nama Uji", "Lulus")])
    output = BytesIO()
    with ZipFile(source) as original, ZipFile(output, "w", ZIP_DEFLATED) as updated:
        for item in original.infolist():
            updated.writestr(item, original.read(item.filename))
        updated.writestr("xl/vbaProject.bin", b"synthetic macro marker")
    output.seek(0)

    with pytest.raises(ImportValidationError, match="macro"):
        parse_sitasi(output, "fixture.xlsx")


def test_import_rejects_oversized_upload():
    with pytest.raises(UploadTooLargeError):
        parse_sitasi(BytesIO(b"x" * (import_service.MAX_UPLOAD_BYTES + 1)), "fixture.xlsx")


def test_sitasi_endpoint_returns_413_for_oversized_file(monkeypatch):
    monkeypatch.setattr(
        sitasi_endpoint,
        "get_settings",
        lambda: SimpleNamespace(max_upload_size_bytes=import_service.MAX_UPLOAD_BYTES),
    )
    upload = UploadFile(
        filename="fixture.xlsx",
        file=BytesIO(b"x" * (import_service.MAX_UPLOAD_BYTES + 1)),
    )

    with pytest.raises(HTTPException) as error:
        sitasi_endpoint.import_sitasi(upload)

    assert error.value.status_code == 413


def test_sitasi_endpoint_hides_internal_storage_error(monkeypatch):
    monkeypatch.setattr(
        sitasi_endpoint,
        "get_settings",
        lambda: SimpleNamespace(max_upload_size_bytes=import_service.MAX_UPLOAD_BYTES),
    )
    monkeypatch.setattr(
        sitasi_endpoint,
        "get_supabase",
        lambda: (_ for _ in ()).throw(RuntimeError("secret student value")),
    )
    upload = UploadFile(
        filename="fixture.xlsx",
        file=sitasi_fixture([sitasi_row("001", "Nama Uji", "Lulus")]),
    )

    with pytest.raises(HTTPException) as error:
        sitasi_endpoint.import_sitasi(upload)

    assert error.value.status_code == 500
    assert error.value.detail == "Import SITASI gagal diproses."
    assert "secret student value" not in error.value.detail


def test_import_success_activates_batch_through_rpc_and_keeps_duplicate_rows():
    supabase = FakeSupabase()
    workbook = sitasi_fixture(
        [
            sitasi_row("001234567", "Nama Uji", "Lulus"),
            sitasi_row("001234567", "Nama Uji", "Gagal"),
        ]
    )

    result = import_master_data(supabase, "SITASI", "fixture.xlsx", workbook)

    assert result["imported_rows"] == 2
    assert len(supabase.rows) == 2
    assert [row["source_row"] for row in supabase.rows] == [2, 3]
    assert supabase.imports[0]["is_active"] is False
    assert supabase.imports[1]["is_active"] is True
    assert supabase.imports[1]["status"] == "COMPLETED"


def test_partial_row_insert_failure_cleans_new_batch_and_preserves_old_active(monkeypatch):
    monkeypatch.setattr(import_service, "BATCH_SIZE", 2)
    supabase = FakeSupabase()
    supabase.fail_row_insert_call = 2
    workbook = sitasi_fixture(
        [
            sitasi_row("001", "Nama Satu", "Lulus"),
            sitasi_row("002", "Nama Dua", "Lulus"),
            sitasi_row("003", "Nama Tiga", "Gagal"),
        ]
    )

    with pytest.raises(RuntimeError, match="batch sebelumnya tetap aktif") as error:
        import_master_data(supabase, "SITASI", "fixture.xlsx", workbook)

    assert "synthetic database detail" not in str(error.value)
    assert supabase.imports[0]["is_active"] is True
    assert supabase.imports[1]["status"] == "FAILED"
    assert supabase.rows == []


def test_rpc_failure_cleans_staging_and_preserves_old_active():
    supabase = FakeSupabase()
    supabase.fail_rpc = True
    workbook = sitasi_fixture([sitasi_row("001", "Nama Uji", "Lulus")])

    with pytest.raises(RuntimeError, match="gagal diaktifkan"):
        import_master_data(supabase, "SITASI", "fixture.xlsx", workbook)

    assert supabase.imports[0]["is_active"] is True
    assert supabase.imports[1]["status"] == "FAILED"
    assert supabase.rows == []


SIMPLIFIED_SITASI_HEADERS = [
    "NIM",
    "Nama",
    "Kode NIM",
    "Program Studi",
    "Sertifikasi",
    "Nilai",
    "Status",
]

SIMPLIFIED_CERTIPORT_HEADERS = ["Nama", "Sertifikasi", "Nilai", "Status"]


def test_parse_sitasi_accepts_simplified_headers():
    workbook = make_xlsx(
        "Master_Data",
        SIMPLIFIED_SITASI_HEADERS,
        [["202331243", "Nama Uji", "31", "S1 Teknik Informatika", "MOS", 850, "Lulus"]],
    )

    parsed = parse_sitasi(workbook, "data_sitasi.xlsx")

    row = parsed.rows[0]
    assert row["nim"] == "202331243"
    assert row["kode_nim"] == "31"
    assert row["program_studi"] == "S1 Teknik Informatika"
    assert row["sertifikasi"] == "MOS"
    assert row["nilai"] == 850
    assert row["status"] == "LULUS"
    assert row["sub_program"] is None
    assert row["fakultas"] is None


def test_parse_certiport_accepts_simplified_headers_and_maps_status():
    workbook = make_xlsx(
        "Results (264) Plus",
        SIMPLIFIED_CERTIPORT_HEADERS,
        [
            ["Nama Uji", "MOS", 900, "Lulus"],
            ["Nama Lain", "MCF", 700, "Tidak Lulus"],
        ],
    )

    parsed = parse_certiport(workbook, "data_certiport.xlsx")

    assert parsed.rows[0]["full_name"] == "Nama Uji"
    assert parsed.rows[0]["full_name_normalized"] == "NAMA UJI"
    assert parsed.rows[0]["program_name"] == "MICROSOFT OFFICE SPECIALIST"
    assert parsed.rows[0]["certification_name"] == "MOS"
    assert parsed.rows[0]["score"] == 900
    assert parsed.rows[0]["result"] == "PASS"
    assert parsed.rows[1]["program_name"] == "MICROSOFT CERTIFIED FUNDAMENTALS"
    assert parsed.rows[1]["result"] == "FAIL"


def test_parse_sitasi_rejects_empty_required_value():
    workbook = make_xlsx(
        "Master_Data",
        SIMPLIFIED_SITASI_HEADERS,
        [["202331243", "", "31", "S1 Teknik Informatika", "MOS", 850, "Lulus"]],
    )

    with pytest.raises(ImportValidationError, match="Nama wajib"):
        parse_sitasi(workbook, "data_sitasi.xlsx")


def test_parse_certiport_rejects_empty_name_and_category():
    empty_name = make_xlsx(
        "Results (264) Plus",
        SIMPLIFIED_CERTIPORT_HEADERS,
        [["", "MOS", 900, "Lulus"]],
    )
    empty_category = make_xlsx(
        "Results (264) Plus",
        SIMPLIFIED_CERTIPORT_HEADERS,
        [["Nama Uji", "", 900, "Lulus"]],
    )

    with pytest.raises(ImportValidationError, match="Nama wajib"):
        parse_certiport(empty_name, "data_certiport.xlsx")
    with pytest.raises(ImportValidationError, match="Sertifikasi wajib"):
        parse_certiport(empty_category, "data_certiport.xlsx")


def test_parse_certiport_preserves_duplicate_rows():
    workbook = make_xlsx(
        "Results (264) Plus",
        SIMPLIFIED_CERTIPORT_HEADERS,
        [
            ["Nama Uji", "MOS", 900, "Lulus"],
            ["Nama Uji", "MOS", 900, "Lulus"],
        ],
    )

    parsed = parse_certiport(workbook, "data_certiport.xlsx")

    assert len(parsed.rows) == 2
    assert [row["source_row"] for row in parsed.rows] == [2, 3]
    assert [row["full_name"] for row in parsed.rows] == ["Nama Uji", "Nama Uji"]
