from __future__ import annotations

from io import BytesIO
import os

os.environ["APP_ENV"] = "test"
os.environ["FRONTEND_URL"] = "http://localhost:3000"
os.environ["SUPABASE_URL"] = "http://127.0.0.1:54321"
os.environ["SUPABASE_SERVICE_ROLE_KEY"] = "dummy-test-key"

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.api.v1.endpoints import sync as sync_endpoint
from app.core.auth import StaffPrincipal, require_staff
from app.main import app
from app.services.matching_service import PERLU_VERIFIKASI, VALID
from app.services.sync_service import generate_sync_export_workbook
from tests.test_sync_service import MockSupabase


@pytest.fixture
def mock_supabase_for_export():
    sb = MockSupabase()
    session_id = "sess-export-1"

    sb.storage["sync_sessions"].append({
        "id": session_id,
        "file_name": "wisuda_2026.xlsx",
        "sitasi_import_id": "sit-batch-uuid",
        "certiport_import_id": "cert-batch-uuid",
        "total_rows": 2,
        "total_valid": 1,
        "total_review": 1,
        "total_not_found": 0,
        "status": "COMPLETED",
        "created_at": "2026-10-03T10:00:00Z",
    })

    sb.storage["sync_results"].extend([
        {
            "id": "res-1",
            "sync_session_id": session_id,
            "source_row": 2,
            "nim": "209900001",
            "nama_wisuda": "MAHASISWA SATU",
            "mos_status": VALID,
            "mcf_status": None,
            "validation_status": VALID,
            "review_note": None,
        },
        {
            "id": "res-2",
            "sync_session_id": session_id,
            "source_row": 3,
            "nim": "209900002",
            "nama_wisuda": "MAHASISWA DUA",
            "mos_status": PERLU_VERIFIKASI,
            "mcf_status": None,
            "validation_status": PERLU_VERIFIKASI,
            "review_note": "Alias DRAFT",
        },
    ])

    return sb, session_id


# ---------------------------------------------------------------------------
# Unit tests for generate_sync_export_workbook
# ---------------------------------------------------------------------------

def test_export_final_fills_only_passed_rows(mock_supabase_for_export):
    sb, session_id = mock_supabase_for_export
    buf = generate_sync_export_workbook(sb, session_id, mode="final")

    wb = openpyxl.load_workbook(buf)
    assert "Wisuda" in wb.sheetnames
    assert "Ringkasan" in wb.sheetnames

    ws = wb["Wisuda"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3  # Header + 2 data rows

    headers = rows[0]
    assert headers[5:8] == ("Sertifikasi MOS", "Sertifikasi MCF/MTCNA", "Title")

    # Baris lulus: terisi.
    row1 = rows[1]
    assert row1[3] == "209900001"
    assert row1[5] == "MOS (Lulus)"
    assert row1[6] in ("", None)
    assert row1[7] == "MOS"

    # Baris tidak lulus: semua kolom target kosong.
    row2 = rows[2]
    assert row2[3] == "209900002"
    assert row2[5] is None
    assert row2[6] is None
    assert row2[7] is None


def test_export_draft_generates_valid_workbook_with_draft_labels(mock_supabase_for_export):
    sb, session_id = mock_supabase_for_export
    buf = generate_sync_export_workbook(sb, session_id, mode="draft")

    wb = openpyxl.load_workbook(buf)
    assert "Wisuda (DRAFT)" in wb.sheetnames
    assert "Ringkasan" in wb.sheetnames

    ws = wb["Wisuda (DRAFT)"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3

    headers = rows[0]
    assert headers[5:8] == ("Sertifikasi MOS", "Sertifikasi MCF/MTCNA", "Title")

    # Baris lulus: berlabel DRAFT.
    row1 = rows[1]
    assert row1[5] == "[DRAFT] MOS (Lulus)"
    assert row1[7] == "[DRAFT] MOS"

    # Baris tidak lulus: kosong.
    row2 = rows[2]
    assert row2[5] is None

    ws_audit = wb["Ringkasan"]
    audit_rows = list(ws_audit.iter_rows(values_only=True))
    assert any("DRAFT - PRATINJAU" in str(r) for r in audit_rows)


# ---------------------------------------------------------------------------
# API tests: GET /api/v1/sync/{session_id}/export
# ---------------------------------------------------------------------------

def test_api_export_accessible_without_auth(monkeypatch, mock_supabase_for_export):
    sb, session_id = mock_supabase_for_export
    monkeypatch.setattr(sync_endpoint, "get_supabase", lambda: sb)
    client = TestClient(app)
    resp = client.get(f"/api/v1/sync/{session_id}/export?mode=draft")
    assert resp.status_code == 200


def test_api_export_final_returns_200_xlsx_stream(monkeypatch, mock_supabase_for_export):
    sb, session_id = mock_supabase_for_export
    monkeypatch.setattr(sync_endpoint, "get_supabase", lambda: sb)

    client = TestClient(app)
    app.dependency_overrides[require_staff] = lambda: StaffPrincipal(
        user_id="staff-user",
        display_name="Staff ITCC",
        role="staff",
    )
    try:
        resp = client.get(f"/api/v1/sync/{session_id}/export?mode=final")
        assert resp.status_code == 200
        assert "spreadsheetml.sheet" in resp.headers["content-type"]
        assert "attachment" in resp.headers["content-disposition"]
        assert "final" in resp.headers["content-disposition"].lower()

        wb = openpyxl.load_workbook(BytesIO(resp.content))
        assert "Wisuda" in wb.sheetnames
    finally:
        app.dependency_overrides.pop(require_staff, None)


def test_api_export_draft_returns_200_xlsx_stream(monkeypatch, mock_supabase_for_export):
    sb, session_id = mock_supabase_for_export
    monkeypatch.setattr(sync_endpoint, "get_supabase", lambda: sb)

    client = TestClient(app)
    app.dependency_overrides[require_staff] = lambda: StaffPrincipal(
        user_id="staff-user",
        display_name="Staff ITCC",
        role="staff",
    )
    try:
        resp = client.get(f"/api/v1/sync/{session_id}/export?mode=draft")
        assert resp.status_code == 200
        assert "spreadsheetml.sheet" in resp.headers["content-type"]
        assert "attachment" in resp.headers["content-disposition"]
        assert "draft" in resp.headers["content-disposition"].lower()

        # Check content is readable as openpyxl workbook
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        assert "Wisuda (DRAFT)" in wb.sheetnames
    finally:
        app.dependency_overrides.pop(require_staff, None)


def test_api_export_not_found_returns_404(monkeypatch, mock_supabase_for_export):
    sb, _ = mock_supabase_for_export
    monkeypatch.setattr(sync_endpoint, "get_supabase", lambda: sb)

    client = TestClient(app)
    app.dependency_overrides[require_staff] = lambda: StaffPrincipal(
        user_id="staff-user",
        display_name="Staff ITCC",
        role="staff",
    )
    try:
        resp = client.get("/api/v1/sync/non-existent-uuid/export?mode=draft")
        assert resp.status_code == 404
    finally:
        app.dependency_overrides.pop(require_staff, None)
