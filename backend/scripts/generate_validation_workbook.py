"""Isi kolom sertifikasi pada workbook Wisuda dari data SITASI dan Certiport.

Sumber (lokal, tidak di-commit):
- data_wisuda.xlsx    sheet "Wisuda"
- data_sitasi.xlsx    sheet "Master_Data"
- data_certiport.xlsx sheet "Results (264) Plus"

Aturan yang disetujui pemilik project (2026-10-23):
- Hanya status LULUS yang diisi.
- Format kolom: "MOS (Lulus)" / "MCF (Lulus)".
- Title: "MOS", "MCF", "MTCNA", atau kombinasinya ("MOS, MTCNA", "MOS, MCF", dst.).
- Sel yang tidak lulus dikosongkan.
- MCF hanya diisi untuk prodi kode 31 dan 32.
- Certiport dicocokkan hanya dengan nama ternormalisasi yang persis sama.

Pipeline membaca berkas memakai parser impor resmi dan mesin matching resmi
(`validate_wisuda_complete`), sehingga hasil script sama dengan backend.

Jalankan dari root repository:
    .\\backend\\venv\\Scripts\\python.exe backend\\scripts\\generate_validation_workbook.py
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path
from typing import Any

BACKEND_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = BACKEND_ROOT.parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from app.services.import_service import parse_certiport, parse_sitasi
from app.services.matching_service import (
    VALID,
    build_wisuda_fill_values,
    status_label,
    validate_wisuda_complete,
)
from app.services.wisuda_validation_service import parse_wisuda_for_sync

WISUDA_SHEET = "Wisuda"
SUMMARY_SHEET = "Ringkasan"
TARGET_COLUMNS = ("Sertifikasi MOS", "Sertifikasi MCF/MTCNA", "Title")


def _read_with_parser(path: Path, parser, *args) -> list[dict]:
    with path.open("rb") as handle:
        parsed = parser(handle, path.name, *args)
    return parsed.rows if hasattr(parsed, "rows") else parsed


def load_sources(wisuda_path: Path, sitasi_path: Path, certiport_path: Path):
    wisuda_rows = _read_with_parser(wisuda_path, parse_wisuda_for_sync, 10 * 1024 * 1024)
    sitasi_rows = _read_with_parser(sitasi_path, parse_sitasi)
    certiport_rows = _read_with_parser(certiport_path, parse_certiport)
    return wisuda_rows, sitasi_rows, certiport_rows


def _header_positions(worksheet) -> tuple[dict[str, int], str, str]:
    for row_number in range(1, min(worksheet.max_row or 0, 10) + 1):
        values = next(
            worksheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True)
        )
        headers = {str(value).strip(): index for index, value in enumerate(values, start=1) if value}
        if "NIM" in headers and "Sertifikasi MOS" in headers:
            mcf_col = "Sertifikasi MCF/MTCNA" if "Sertifikasi MCF/MTCNA" in headers else "Sertifikasi MCF"
            title_col = "Title" if "Title" in headers else "Title Microsoft"
            return headers, mcf_col, title_col
    raise ValueError("Header sheet Wisuda tidak dikenali (NIM/Sertifikasi MOS tidak ditemukan).")


def fill_wisuda_workbook(
    source_path: Path,
    output_path: Path,
    results: list[dict],
) -> Counter:
    """Salin workbook Wisuda, isi kolom target, lalu simpan ke berkas baru."""
    workbook = load_workbook(source_path)
    if WISUDA_SHEET not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"Sheet wajib '{WISUDA_SHEET}' tidak ditemukan.")

    worksheet = workbook[WISUDA_SHEET]
    positions, mcf_col, title_col = _header_positions(worksheet)
    target_cols = ("Sertifikasi MOS", mcf_col, title_col)
    missing = [column for column in target_cols if column not in positions]
    if missing:
        workbook.close()
        raise ValueError(f"Kolom target tidak ditemukan: {', '.join(missing)}")

    counts: Counter = Counter()
    for result in results:
        row_number = result.get("source_row")
        if not isinstance(row_number, int):
            continue
        prog_summary = result.get("program_summary", {})
        mcf_prog = prog_summary.get("MCF_COMBINED") or prog_summary.get("MCF") or prog_summary.get("MTCNA")
        fill = build_wisuda_fill_values(result.get("mos_status"), result.get("mcf_status"), mcf_prog)
        for column, value in zip(target_cols, (fill["mos"], fill["mcf"], fill["title"])):
            worksheet.cell(row=row_number, column=positions[column]).value = value or None

        counts["total"] += 1
        if result.get("mos_status") == VALID:
            counts["mos"] += 1
        if result.get("mcf_status") == VALID:
            counts["mcf"] += 1
        if result.get("mos_status") == VALID and result.get("mcf_status") == VALID:
            counts["both"] += 1
        if result.get("mos_status") != VALID and result.get("mcf_status") != VALID:
            counts["none"] += 1

    if SUMMARY_SHEET in workbook.sheetnames:
        del workbook[SUMMARY_SHEET]
    summary = workbook.create_sheet(SUMMARY_SHEET)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 70
    summary.append(["PARAMETER", "NILAI"])
    summary.cell(row=1, column=1).font = Font(bold=True)
    summary.cell(row=1, column=2).font = Font(bold=True)
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor="F1F5F9")

    rows: list[tuple[str, Any]] = [
        ("Sumber Wisuda", source_path.name),
        ("Aturan isi", "Hanya status LULUS yang diisi; sel lain kosong."),
        ("Format MOS", "MOS (Lulus)"),
        ("Format MCF/MTCNA", "MCF (Lulus) atau MTCNA (Lulus) untuk prodi 31 dan 32"),
        ("Format Title", "MOS, MCF, MTCNA, atau kombinasinya"),
        ("Total mahasiswa", counts["total"]),
        ("Lulus MOS", counts["mos"]),
        ("Lulus MCF", counts["mcf"]),
        ("Lulus MOS dan MCF", counts["both"]),
        ("Tidak ada isian", counts["none"]),
    ]
    for label, value in rows:
        summary.append([label, value])

    workbook.save(output_path)
    workbook.close()
    return counts


def main() -> int:
    parser = argparse.ArgumentParser(description="Isi kolom sertifikasi workbook Wisuda.")
    parser.add_argument("--wisuda", type=Path, default=REPO_ROOT / "data_wisuda.xlsx")
    parser.add_argument("--sitasi", type=Path, default=REPO_ROOT / "data_sitasi.xlsx")
    parser.add_argument("--certiport", type=Path, default=REPO_ROOT / "data_certiport.xlsx")
    parser.add_argument("--output", type=Path, default=REPO_ROOT / "hasil_wisuda_terisi.xlsx")
    args = parser.parse_args()

    for path in (args.wisuda, args.sitasi, args.certiport):
        if not path.exists():
            print(f"File sumber tidak ditemukan: {path}", file=sys.stderr)
            return 1

    print("Membaca file sumber...")
    wisuda_rows, sitasi_rows, certiport_rows = load_sources(args.wisuda, args.sitasi, args.certiport)
    print(f"  Wisuda: {len(wisuda_rows)} baris")
    print(f"  SITASI: {len(sitasi_rows)} baris")
    print(f"  Certiport: {len(certiport_rows)} baris")

    print("Menjalankan validasi...")
    results = validate_wisuda_complete(
        wisuda_rows,
        sitasi_rows,
        certiport_rows,
        sitasi_import_id=f"FILE:{args.sitasi.name}",
        certiport_import_id=f"FILE:{args.certiport.name}",
        certiport_file_name=args.certiport.name,
    )

    print("Menulis workbook hasil...")
    counts = fill_wisuda_workbook(args.wisuda, args.output, results)

    assert counts["total"] == len(wisuda_rows), "Jumlah baris hasil harus sama dengan baris Wisuda."
    print(f"Selesai: {args.output.resolve()}")
    print(
        "Ringkasan: "
        f"total={counts['total']}, MOS lulus={counts['mos']}, MCF lulus={counts['mcf']}, "
        f"MOS+MCF={counts['both']}, tanpa isian={counts['none']}"
    )
    labels = Counter(status_label(result.get("validation_status")) for result in results)
    print("Status validasi: " + ", ".join(f"{label}={count}" for label, count in sorted(labels.items())))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
