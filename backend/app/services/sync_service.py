from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, BinaryIO
import copy
import uuid

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from supabase import Client

from app.services.import_service import ImportValidationError, get_active_import_id
from app.services.matching_service import (
    DATA_TIDAK_DITEMUKAN,
    PENDING,
    PERLU_VERIFIKASI,
    TIDAK_LULUS,
    VALID,
    fetch_active_certiport_rows,
    fetch_active_sitasi_rows,
    build_wisuda_fill_values,
    validate_wisuda_complete,
)
from app.services.wisuda_validation_service import parse_wisuda_for_sync
from app.core.prodi import is_mcf_relevant, resolve_kode_prodi, resolve_program_studi


class SyncBatchMissingError(ValueError):
    """Batch master SITASI atau Certiport aktif belum tersedia."""


class SyncSessionNotFoundError(ValueError):
    """Sesi sinkronisasi tidak ditemukan."""


class SyncExecutionError(RuntimeError):
    """Kegagalan selama proses sinkronisasi."""


class ReviewPolicyBlockedError(PermissionError):
    """ADR-009: Kebijakan penunjukan reviewer belum disahkan."""


class ReviewValidationError(ValueError):
    """Input keputusan atau catatan review tidak valid."""


class ExportBlockedByPolicyError(PermissionError):
    """ADR-003/005/006/010/011: Ekspor final diblokir sebelum pengesahan resmi ITCC."""


ALLOWED_REVIEW_DECISIONS = ("KONFIRMASI_COCOK", "TIDAK_COCOK", "TUNDA_VERIFIKASI")
ADR_009_REVIEWER_POLICY_APPROVED = False  # ADR-009 berstatus TERBUKA per docs/14_DECISIONS.md
# Disetujui pemilik project 2026-10-23: isi kolom hanya untuk status LULUS,
# format "MOS (Lulus)"/"MCF (Lulus)", Title "MOS"/"MCF", berkas keluaran baru.
ADR_OUTPUT_MAPPING_APPROVED = True


def _get_import_file_name(supabase: Client, import_id: str | None) -> str | None:
    """Ambil nama file batch aktif untuk jejak audit kandidat."""
    if not import_id:
        return None
    try:
        response = (
            supabase.table("data_imports")
            .select("id,file_name")
            .eq("id", import_id)
            .limit(1)
            .execute()
        )
        if response.data:
            return response.data[0].get("file_name")
    except Exception:
        pass
    return None


def create_and_execute_sync(
    supabase: Client,
    file_obj: BinaryIO,
    file_name: str,
    max_upload_bytes: int,
    user_id: str | None = None,
) -> dict:
    """Eksekusi sinkronisasi wisuda end-to-end dengan persistensi sesi, hasil, dan bukti.

    Alur:
    1. Parse dan validasi baris wisuda (menolak formula, duplikat, dsb.)
    2. Cek batch SITASI dan Certiport aktif. Catat snapshot ID batch pada sesi.
    3. Simpan sesi awal dengan status PROCESSING.
    4. Ambil master data batch aktif terkait.
    5. Jalankan validasi silang lengkap (validate_wisuda_complete).
    6. Persistensi hasil ke sync_results dan detail bukti ke sync_evidence.
    7. Perbarui status sesi menjadi COMPLETED beserta ringkasan total.
    8. Jika gagal di tengah, status sesi diubah menjadi FAILED beserta pesan error aman.
    """
    wisuda_rows = parse_wisuda_for_sync(file_obj, file_name, max_upload_bytes)

    sitasi_import_id = get_active_import_id(supabase, "SITASI")
    certiport_import_id = get_active_import_id(supabase, "CERTIPORT")

    if not sitasi_import_id or not certiport_import_id:
        missing = []
        if not sitasi_import_id:
            missing.append("SITASI")
        if not certiport_import_id:
            missing.append("CERTIPORT")
        raise SyncBatchMissingError(
            f"Batch master aktif belum lengkap ({', '.join(missing)}). "
            f"Unggah dan aktifkan batch master terlebih dahulu."
        )

    session_id = str(uuid.uuid4())
    now_iso = datetime.now(timezone.utc).isoformat()

    session_data = {
        "id": session_id,
        "file_name": file_name,
        "sitasi_import_id": sitasi_import_id,
        "certiport_import_id": certiport_import_id,
        "total_rows": len(wisuda_rows),
        "total_valid": 0,
        "total_review": 0,
        "total_not_found": 0,
        "status": "PROCESSING",
        "created_by": user_id,
        "created_at": now_iso,
    }

    supabase.table("sync_sessions").insert(session_data).execute()

    try:
        _, sitasi_rows = fetch_active_sitasi_rows(supabase)
        _, certiport_rows = fetch_active_certiport_rows(supabase)

        validated_rows = validate_wisuda_complete(
            wisuda_rows,
            sitasi_rows,
            certiport_rows,
            sitasi_import_id=sitasi_import_id,
            certiport_import_id=certiport_import_id,
            certiport_file_name=_get_import_file_name(supabase, certiport_import_id),
        )

        results_to_insert: list[dict] = []
        evidence_to_insert: list[dict] = []

        total_valid = 0
        total_review = 0
        total_not_found = 0

        for r in validated_rows:
            result_id = str(uuid.uuid4())
            v_status = r.get("validation_status", PERLU_VERIFIKASI)
            needs_review = bool(r.get("needs_review", False))

            if v_status == VALID:
                total_valid += 1
            elif v_status == DATA_TIDAK_DITEMUKAN:
                total_not_found += 1
            else:
                total_review += 1

            review_flags = r.get("review_flags", [])
            review_note = "; ".join(review_flags) if review_flags else None

            prog_summary = r.get("program_summary", {})
            mos_programs = prog_summary.get("MOS")
            mcf_programs = (
                prog_summary.get("MCF_COMBINED")
                or prog_summary.get("MCF")
                or prog_summary.get("MTCNA")
            )

            nama_sitasi = None
            exams = r.get("exams", [])
            for exam in exams:
                for att in exam.get("attempts", []):
                    if att.get("nama"):
                        nama_sitasi = att["nama"]
                        break
                if nama_sitasi:
                    break

            nim_val = str(r["nim"])
            kode_prodi_val = resolve_kode_prodi(nim_val, r.get("kode_prodi"))
            prodi_val = resolve_program_studi(kode_prodi_val, r.get("program_studi"))

            certiport_match = r.get("certiport_match", {}) or {}
            if r.get("name_match"):
                name_match_type = r.get("name_match")
            elif certiport_match.get("status") == "MATCHED":
                name_match_type = "CERTIPORT_EXACT"
            elif certiport_match.get("status") == PERLU_VERIFIKASI:
                name_match_type = "CERTIPORT_CANDIDATE_ONLY"
            else:
                name_match_type = "TIDAK_ADA_IDENTITAS"

            results_to_insert.append({
                "id": result_id,
                "sync_session_id": session_id,
                "source_row": r["source_row"],
                "nim": nim_val,
                "nama_wisuda": str(r["nama"]),
                "kode_prodi": kode_prodi_val,
                "program_studi": prodi_val,
                "nama_sitasi": nama_sitasi,
                "name_match_type": name_match_type,
                "mos_status": r.get("mos_status"),
                "mcf_status": r.get("mcf_status"),
                "mos_programs": mos_programs,
                "mcf_programs": mcf_programs,
                "validation_status": v_status,
                "needs_review": needs_review,
                "review_note": review_note,
                "created_at": now_iso,
            })

            for exam in exams:
                cat = exam.get("sertifikasi")
                sub = exam.get("sub_program") or cat
                cv = exam.get("certiport_validation", {})
                e_status = cv.get("combined_status", exam.get("exam_status"))
                e_reason = cv.get("reason", exam.get("reason", ""))

                sitasi_ev = exam.get("lulus_evidence") or exam.get("selected_attempt") or (exam.get("attempts", [{}])[0] if exam.get("attempts") else {})
                sit_row_id = sitasi_ev.get("id")

                cert_ev_list = cv.get("certiport_evidence", [])
                cert_row_id = cert_ev_list[0].get("id") if cert_ev_list else None

                sit_score = sitasi_ev.get("nilai")
                cert_score = cert_ev_list[0].get("score") if cert_ev_list else None
                score = sit_score if sit_score is not None else cert_score

                evidence_to_insert.append({
                    "id": str(uuid.uuid4()),
                    "result_id": result_id,
                    "category": cat,
                    "exam": sub,
                    "sitasi_row_id": sit_row_id,
                    "certiport_row_id": cert_row_id,
                    "match_method": "NIM_EXACT+NAME_CERTIPORT",
                    "match_score": score,
                    "evidence_status": e_status,
                    "reason": e_reason,
                    "created_at": now_iso,
                })

        batch_size = 500
        for i in range(0, len(results_to_insert), batch_size):
            chunk = results_to_insert[i : i + batch_size]
            supabase.table("sync_results").insert(chunk).execute()

        for i in range(0, len(evidence_to_insert), batch_size):
            chunk = evidence_to_insert[i : i + batch_size]
            supabase.table("sync_evidence").insert(chunk).execute()

        completed_at = datetime.now(timezone.utc).isoformat()
        supabase.table("sync_sessions").update({
            "status": "COMPLETED",
            "total_valid": total_valid,
            "total_review": total_review,
            "total_not_found": total_not_found,
            "completed_at": completed_at,
        }).eq("id", session_id).execute()

        return {
            "session_id": session_id,
            "status": "COMPLETED",
            "file_name": file_name,
            "total_rows": len(wisuda_rows),
            "total_valid": total_valid,
            "total_review": total_review,
            "total_not_found": total_not_found,
            "sitasi_import_id": sitasi_import_id,
            "certiport_import_id": certiport_import_id,
            "completed_at": completed_at,
        }

    except Exception as exc:
        try:
            supabase.table("sync_sessions").update({
                "status": "FAILED",
                "error_message": "Terjadi kesalahan saat memproses data sinkronisasi.",
            }).eq("id", session_id).execute()
        except Exception:
            pass
        raise SyncExecutionError("Proses sinkronisasi gagal diproses.") from exc


def list_sync_sessions(
    supabase: Client,
    page: int = 1,
    limit: int = 50,
) -> dict:
    """Daftar riwayat sesi sinkronisasi dengan paginasi."""
    limit = max(1, min(limit, 100))
    offset = max(0, (page - 1) * limit)

    response = (
        supabase.table("sync_sessions")
        .select("*", count="exact")
        .order("created_at", desc=True)
        .range(offset, offset + limit - 1)
        .execute()
    )
    items = response.data or []
    total = response.count if response.count is not None else len(items)

    return {
        "items": items,
        "page": page,
        "limit": limit,
        "total": total,
    }


def get_sync_session(
    supabase: Client,
    session_id: str,
) -> dict:
    """Ambil ringkasan satu sesi sinkronisasi berdasarkan ID."""
    response = (
        supabase.table("sync_sessions")
        .select("*")
        .eq("id", session_id)
        .limit(1)
        .execute()
    )
    if not response.data:
        raise SyncSessionNotFoundError(f"Sesi sinkronisasi {session_id} tidak ditemukan.")
    return response.data[0]


def _row_has_program(r: dict, prog_code: str) -> bool:
    """Periksa apakah baris hasil sinkronisasi memuat program sertifikasi tertentu."""
    code = prog_code.strip().upper()
    if code == "MOS":
        mos_stat = str(r.get("mos_status") or "").upper()
        if mos_stat and mos_stat != "DATA_TIDAK_DITEMUKAN":
            return True
        mos_p = r.get("mos_programs")
        if isinstance(mos_p, dict) and (mos_p.get("exam_details") or mos_p.get("scores") or mos_p.get("score") is not None):
            return True
        return False
    elif code == "MTCNA":
        mcf_p = r.get("mcf_programs")
        if isinstance(mcf_p, dict):
            for ed in mcf_p.get("exam_details") or []:
                cert = str(ed.get("sertifikasi") or ed.get("sub_program") or "").upper()
                if "MTCNA" in cert:
                    return True
        return False
    elif code == "MCF":
        mcf_p = r.get("mcf_programs")
        if isinstance(mcf_p, dict):
            for ed in mcf_p.get("exam_details") or []:
                cert = str(ed.get("sertifikasi") or ed.get("sub_program") or "").upper()
                if any(k in cert for k in ("MCF", "AZURE", "AI-900", "AZ-900", "SC-900", "DP-900", "PL-900", "MS-900")):
                    return True
            if not any("MTCNA" in str(ed.get("sertifikasi") or ed.get("sub_program") or "").upper() for ed in (mcf_p.get("exam_details") or [])):
                if r.get("mcf_status") and r.get("mcf_status") != "DATA_TIDAK_DITEMUKAN":
                    return True
        elif r.get("mcf_status") and r.get("mcf_status") != "DATA_TIDAK_DITEMUKAN":
            return True
        return False
    return True


def list_sync_results(
    supabase: Client,
    session_id: str,
    page: int = 1,
    limit: int = 50,
    status: str | None = None,
    search: str | None = None,
    needs_review: bool | None = None,
    program: str | None = None,
) -> dict:
    """Ambil daftar hasil per baris dengan paginasi dan filter status serta program."""
    get_sync_session(supabase, session_id)

    limit = max(1, min(limit, 100))
    offset = max(0, (page - 1) * limit)

    query = (
        supabase.table("sync_results")
        .select("*", count="exact")
        .eq("sync_session_id", session_id)
    )

    if status:
        status_clean = status.strip().upper()
        if status_clean == PERLU_VERIFIKASI:
            query = query.in_("validation_status", [PERLU_VERIFIKASI, TIDAK_LULUS])
        else:
            query = query.eq("validation_status", status_clean)

    if needs_review is not None:
        query = query.eq("needs_review", needs_review)

    if search:
        s = search.strip()
        query = query.or_(f"nim.ilike.%{s}%,nama_wisuda.ilike.%{s}%")

    if program:
        # Ambil seluruh record yang cocok dengan filter dasar, lalu saring berdasarkan program
        query = query.order("source_row", desc=False)
        response = query.execute()
        all_items = response.data or []
        filtered_items = [it for it in all_items if _row_has_program(it, program)]
        total = len(filtered_items)
        items = filtered_items[offset : offset + limit]
    else:
        query = query.order("source_row", desc=False).range(offset, offset + limit - 1)
        response = query.execute()
        items = response.data or []
        total = response.count if response.count is not None else len(items)

    for it in items:
        if not it.get("kode_prodi"):
            it["kode_prodi"] = resolve_kode_prodi(it.get("nim"))
        if not it.get("program_studi"):
            it["program_studi"] = resolve_program_studi(it.get("kode_prodi"))

    return {
        "items": items,
        "page": page,
        "limit": limit,
        "total": total,
    }


def _evidence_pass_rank(ev: dict) -> int:
    """Prioritas tampil saat satu (kategori, sub-ujian) punya bukti ganda.

    0 = ada bukti lulus (ditampilkan), 1 = sisanya (disembunyikan).
    Tidak mengubah status — hanya memilih tampilan.
    """
    if str(ev.get("evidence_status", "")).upper() == VALID:
        return 0
    if str(ev.get("certiport_status", "")).upper() == "PASS":
        return 0
    if str(ev.get("sitasi_status", "")).upper() == "LULUS":
        return 0
    return 1


def _normalize_master_certiport_result(value: Any) -> str | None:
    """Normalkan kolom `result` master Certiport menjadi PASS/FAIL."""
    text = str(value or "").strip().upper()
    if text in {"PASS", "PASSED", "LULUS"}:
        return "PASS"
    if text in {"FAIL", "FAILED", "GAGAL", "TIDAK LULUS"}:
        return "FAIL"
    return None


def dedupe_evidence_for_display(evidence: list[dict]) -> list[dict]:
    """Kelompokkan bukti ganda per (kategori, sub-ujian); tampilkan yang lulus.

    Kandidat fuzzy Certiport ganda (mis. dua baris Word mirip) menghasilkan
    beberapa baris evidence untuk sub-ujian yang sama. Baris yang
    disembunyikan tetap tersimpan di database; jumlahnya dicatat di
    `duplicate_count` agar transparan bagi reviewer.
    """
    grouped: dict[tuple[str, str], list[dict]] = {}
    for ev in evidence:
        key = (str(ev.get("category") or "").upper(), str(ev.get("exam") or "").upper())
        grouped.setdefault(key, []).append(ev)

    deduped: list[dict] = []
    for key in grouped:
        group = grouped[key]
        if len(group) == 1:
            shown = dict(group[0])
            shown["duplicate_count"] = 0
            shown["duplicate_ids"] = []
        else:
            ordered = sorted(group, key=_evidence_pass_rank)
            shown = dict(ordered[0])
            shown["duplicate_count"] = len(group) - 1
            shown["duplicate_ids"] = [g.get("id") for g in ordered[1:]]
        deduped.append(shown)
    return deduped


def attach_evidence_source_refs(supabase: Client, evidence: list[dict]) -> None:
    """Lengkapi tiap evidence dengan baris xlsx dan nama file sumber.

    `sitasi_row_id`/`certiport_row_id` adalah id database, bukan baris xlsx.
    Resolve ke `source_row` + `file_name` batch import agar reviewer tahu
    nama tersebut ada di baris mana pada file apa
    (`data_sitasi.xlsx` / `data_certiport.xlsx`). Gagal resolve tidak
    menggagalkan detail — field dibiarkan None.
    """
    try:
        sitasi_ids = sorted({ev.get("sitasi_row_id") for ev in evidence if ev.get("sitasi_row_id") is not None})
        certiport_ids = sorted({ev.get("certiport_row_id") for ev in evidence if ev.get("certiport_row_id") is not None})

        master_lookup: dict[str, dict] = {}
        if sitasi_ids:
            resp = supabase.table("sitasi_data").select("id,source_row,import_id,status,nilai").in_("id", sitasi_ids).execute()
            for row in resp.data or []:
                master_lookup[f"S:{row.get('id')}"] = row
        if certiport_ids:
            resp = supabase.table("certiport_data").select("id,source_row,import_id,score,result").in_("id", certiport_ids).execute()
            for row in resp.data or []:
                master_lookup[f"C:{row.get('id')}"] = row

        file_names: dict[str, str] = {}
        import_ids = sorted({r.get("import_id") for r in master_lookup.values() if r.get("import_id")})
        if import_ids:
            resp = supabase.table("data_imports").select("id,file_name").in_("id", import_ids).execute()
            for row in resp.data or []:
                if row.get("id") is not None and row.get("file_name"):
                    file_names[str(row["id"])] = str(row["file_name"])

        for ev in evidence:
            sit = master_lookup.get(f"S:{ev.get('sitasi_row_id')}", {}) if ev.get("sitasi_row_id") is not None else {}
            crt = master_lookup.get(f"C:{ev.get('certiport_row_id')}", {}) if ev.get("certiport_row_id") is not None else {}
            ev["sitasi_source_row"] = sit.get("source_row")
            ev["sitasi_file_name"] = file_names.get(str(sit["import_id"])) if sit.get("import_id") else None
            ev["certiport_source_row"] = crt.get("source_row")
            ev["certiport_file_name"] = file_names.get(str(crt["import_id"])) if crt.get("import_id") else None
            # Koreksi status tampilan dari rekaman master (fallback lama menebak
            # PASS dari evidence_status sehingga record FAIL ikut berlabel PASS).
            master_result = _normalize_master_certiport_result(crt.get("result"))
            if master_result:
                ev["certiport_status"] = master_result
            if not ev.get("sitasi_status") and sit.get("status"):
                ev["sitasi_status"] = sit["status"]
            if ev.get("certiport_score") is None and crt.get("score") is not None:
                ev["certiport_score"] = crt["score"]
            if ev.get("sitasi_score") is None and sit.get("nilai") is not None:
                ev["sitasi_score"] = sit["nilai"]
    except Exception:
        for ev in evidence:
            ev.setdefault("sitasi_source_row", None)
            ev.setdefault("sitasi_file_name", None)
            ev.setdefault("certiport_source_row", None)
            ev.setdefault("certiport_file_name", None)


def get_sync_result_detail(
    supabase: Client,
    session_id: str,
    result_id: str,
) -> dict:
    """Ambil detail satu baris hasil sinkronisasi beserta seluruh evidence sumber."""
    res_resp = (
        supabase.table("sync_results")
        .select("*")
        .eq("id", result_id)
        .eq("sync_session_id", session_id)
        .limit(1)
        .execute()
    )
    if not res_resp.data:
        raise SyncSessionNotFoundError("Hasil sinkronisasi tidak ditemukan.")

    result_row = res_resp.data[0]

    ev_resp = (
        supabase.table("sync_evidence")
        .select("*")
        .eq("result_id", result_id)
        .order("created_at")
        .execute()
    )
    raw_evidence = ev_resp.data or []

    programs_map = {
        "MOS": result_row.get("mos_programs") or {},
        "MCF": result_row.get("mcf_programs") or {},
    }

    enriched_evidence = []
    for ev in raw_evidence:
        cat = (ev.get("category") or "").upper()
        exam_name = ev.get("exam") or ""
        prog_data = programs_map.get(cat, {})
        exam_details = prog_data.get("exam_details", []) if isinstance(prog_data, dict) else []

        matched_detail = next(
            (d for d in exam_details if str(d.get("sub_program", "")).upper() == str(exam_name).upper()),
            None
        )

        ev_copy = dict(ev)
        has_sitasi_row = ev.get("sitasi_row_id") is not None
        has_cert_row = ev.get("certiport_row_id") is not None

        if matched_detail:
            ev_copy["sitasi_name"] = matched_detail.get("sitasi_name") if has_sitasi_row else None
            ev_copy["sitasi_score"] = matched_detail.get("sitasi_score") if has_sitasi_row else None
            ev_copy["sitasi_status"] = matched_detail.get("sitasi_status") if has_sitasi_row else None

            ev_copy["certiport_name"] = matched_detail.get("certiport_name") if has_cert_row else None
            ev_copy["certiport_score"] = matched_detail.get("certiport_score") if has_cert_row else None
            ev_copy["certiport_status"] = matched_detail.get("certiport_status") if has_cert_row else None
        else:
            ev_copy["sitasi_name"] = result_row.get("nama_sitasi") if has_sitasi_row else None
            ev_copy["sitasi_score"] = None
            ev_copy["sitasi_status"] = None
            ev_copy["certiport_name"] = None
            ev_copy["certiport_score"] = None
            ev_copy["certiport_status"] = None

        # Fallback atribuasi skor yang tepat untuk rekaman tunggal / sesi yang sudah ada
        if has_cert_row and not has_sitasi_row:
            ev_copy["sitasi_name"] = None
            ev_copy["sitasi_score"] = None
            ev_copy["sitasi_status"] = None
            if ev_copy.get("certiport_score") is None:
                ev_copy["certiport_score"] = ev.get("match_score")
            if not ev_copy.get("certiport_name"):
                ev_copy["certiport_name"] = result_row.get("nama_wisuda")
            if not ev_copy.get("certiport_status"):
                ev_status_upper = str(ev.get("evidence_status", "")).upper()
                ev_copy["certiport_status"] = "FAIL" if "FAIL" in ev_status_upper else "PASS"
        elif has_sitasi_row and not has_cert_row:
            ev_copy["certiport_name"] = None
            ev_copy["certiport_score"] = None
            ev_copy["certiport_status"] = None
            if ev_copy.get("sitasi_score") is None:
                ev_copy["sitasi_score"] = ev.get("match_score")
            if not ev_copy.get("sitasi_name"):
                ev_copy["sitasi_name"] = result_row.get("nama_sitasi") or result_row.get("nama_wisuda")
        elif has_sitasi_row and has_cert_row:
            if not ev_copy.get("sitasi_name"):
                ev_copy["sitasi_name"] = result_row.get("nama_sitasi") or result_row.get("nama_wisuda")
            if not ev_copy.get("certiport_name"):
                ev_copy["certiport_name"] = result_row.get("nama_wisuda")
            if ev_copy.get("certiport_score") is None and ev.get("match_score") is not None:
                ev_copy["certiport_score"] = ev.get("match_score")
            if ev_copy.get("sitasi_score") is None and ev.get("match_score") is not None:
                ev_copy["sitasi_score"] = ev.get("match_score")

        enriched_evidence.append(ev_copy)

    attach_evidence_source_refs(supabase, enriched_evidence)
    result_row["evidence"] = dedupe_evidence_for_display(enriched_evidence)

    audit_resp = (
        supabase.table("review_audit")
        .select("*")
        .eq("result_id", result_id)
        .order("created_at", desc=True)
        .execute()
    )
    result_row["review_history"] = audit_resp.data or []

    kp = resolve_kode_prodi(result_row.get("nim"), result_row.get("kode_prodi"))
    result_row["kode_prodi"] = kp
    result_row["program_studi"] = resolve_program_studi(kp, result_row.get("program_studi"))

    return result_row


def get_review_history(supabase: Client, result_id: str) -> list[dict]:
    """Ambil riwayat audit keputusan review untuk satu hasil sinkronisasi."""
    response = (
        supabase.table("review_audit")
        .select("*")
        .eq("result_id", result_id)
        .order("created_at", desc=True)
        .execute()
    )
    return response.data or []


def submit_review(
    supabase: Client,
    session_id: str,
    result_id: str,
    decision: str,
    note: str,
    target_program: str | None = None,
    user_id: str | None = None,
    user_name: str | None = None,
    user_role: str = "admin",
) -> dict:
    """Simpan keputusan review manual petugas dengan audit trail lengkap.

    Aturan:
    - Keputusan harus salah satu dari: KONFIRMASI_COCOK, TIDAK_COCOK, TUNDA_VERIFIKASI.
    - Catatan (note) wajib diisi.
    - Otorisasi reviewer mengikuti ADR-009:
      Jika kebijakan belum disahkan dan user bukan admin, KONFIRMASI_COCOK ditolak
      (mode read-only / tunda verifikasi) untuk mencegah status VALID tanpa izin institusi.
    - Review bersifat append-only di review_audit; bukti asli di sync_evidence TIDAK dihapus.
    - Review ganda mencatat status sebelumnya dan status baru secara berurutan.
    - Memperbarui validation_status, mos_status, mcf_status, mos_programs, mcf_programs,
      serta status sub-ujian evidence yang relevan sehingga tampilan tabel, modal, dan berkas ekspor sinkron.
    """
    decision_clean = decision.strip().upper() if decision else ""
    if decision_clean not in ALLOWED_REVIEW_DECISIONS:
        raise ReviewValidationError(
            f"Keputusan '{decision}' tidak valid. Pilihan: {', '.join(ALLOWED_REVIEW_DECISIONS)}."
        )

    note_clean = note.strip() if note else ""
    if not note_clean:
        raise ReviewValidationError("Catatan review wajib diisi.")

    # ADR-009 check: Kebijakan penunjukan reviewer belum disahkan
    if decision_clean == "KONFIRMASI_COCOK":
        if user_role != "admin" and not ADR_009_REVIEWER_POLICY_APPROVED:
            raise ReviewPolicyBlockedError(
                "Kebijakan penunjukan reviewer (ADR-009) belum disahkan institusi. "
                "Konfirmasi final dibatasi bagi penanggung jawab resmi/admin. "
                "Gunakan 'TUNDA_VERIFIKASI' atau 'TIDAK_COCOK'."
            )

    res_resp = (
        supabase.table("sync_results")
        .select("*")
        .eq("id", result_id)
        .eq("sync_session_id", session_id)
        .limit(1)
        .execute()
    )
    if not res_resp.data:
        raise SyncSessionNotFoundError("Hasil sinkronisasi tidak ditemukan.")

    current_result = res_resp.data[0]
    previous_status = current_result.get("validation_status", PERLU_VERIFIKASI)

    target_clean = (target_program or "ALL").strip().upper()
    current_mos = current_result.get("mos_status")
    current_mcf = current_result.get("mcf_status")

    mos_prog = (
        copy.deepcopy(current_result.get("mos_programs"))
        if isinstance(current_result.get("mos_programs"), dict)
        else None
    )
    mcf_prog = (
        copy.deepcopy(current_result.get("mcf_programs"))
        if isinstance(current_result.get("mcf_programs"), dict)
        else None
    )

    kode_prodi = current_result.get("kode_prodi") or resolve_kode_prodi(current_result.get("nim"))
    mcf_rel = is_mcf_relevant(kode_prodi)

    new_mos_status = current_mos
    new_mcf_status = current_mcf

    if decision_clean == "KONFIRMASI_COCOK":
        if target_clean in ("ALL", "MOS") and current_mos == PERLU_VERIFIKASI:
            new_mos_status = VALID
            if mos_prog:
                mos_prog["status"] = VALID
                mos_prog["needs_review"] = False
        if target_clean in ("ALL", "MCF") and current_mcf == PERLU_VERIFIKASI:
            new_mcf_status = VALID
            if mcf_prog:
                mcf_prog["status"] = VALID
                mcf_prog["needs_review"] = False

    elif decision_clean == "TIDAK_COCOK":
        if target_clean in ("ALL", "MOS") and current_mos == PERLU_VERIFIKASI:
            new_mos_status = TIDAK_LULUS
            if mos_prog:
                mos_prog["status"] = TIDAK_LULUS
                mos_prog["needs_review"] = False
        if target_clean in ("ALL", "MCF") and current_mcf == PERLU_VERIFIKASI:
            new_mcf_status = TIDAK_LULUS
            if mcf_prog:
                mcf_prog["status"] = TIDAK_LULUS
                mcf_prog["needs_review"] = False

    elif decision_clean == "TUNDA_VERIFIKASI":
        # Tunda verifikasi tetap mempertahankan status PERLU_VERIFIKASI
        pass

    # Evaluasi status keseluruhan (validation_status)
    relevant_statuses = [new_mos_status] if new_mos_status else []
    if mcf_rel and new_mcf_status:
        relevant_statuses.append(new_mcf_status)

    if PERLU_VERIFIKASI in relevant_statuses or decision_clean == "TUNDA_VERIFIKASI":
        new_status = PERLU_VERIFIKASI
        needs_review = True
    elif decision_clean == "KONFIRMASI_COCOK":
        new_status = VALID
        needs_review = False
    elif decision_clean == "TIDAK_COCOK":
        if any(s == VALID for s in relevant_statuses):
            new_status = VALID
            needs_review = False
        else:
            new_status = TIDAK_LULUS
            needs_review = False
    else:
        new_status = PERLU_VERIFIKASI
        needs_review = True

    audit_id = str(uuid.uuid4())
    now_iso = datetime.now(timezone.utc).isoformat()

    audit_entry = {
        "id": audit_id,
        "result_id": result_id,
        "reviewed_by": user_id,
        "reviewer_name": user_name or "Petugas",
        "decision": decision_clean,
        "note": note_clean,
        "previous_status": previous_status,
        "new_status": new_status,
        "created_at": now_iso,
    }
    supabase.table("review_audit").insert(audit_entry).execute()

    update_payload: dict[str, Any] = {
        "validation_status": new_status,
        "mos_status": new_mos_status,
        "mcf_status": new_mcf_status,
        "needs_review": needs_review,
        "review_note": f"[{decision_clean}] {note_clean}",
    }
    if mos_prog is not None:
        update_payload["mos_programs"] = mos_prog
    if mcf_prog is not None:
        update_payload["mcf_programs"] = mcf_prog

    supabase.table("sync_results").update(update_payload).eq("id", result_id).execute()

    # Perbarui evidence_status pada sync_evidence untuk sub-ujian yang bersesuaian
    if decision_clean in ("KONFIRMASI_COCOK", "TIDAK_COCOK"):
        evidence_new_status = VALID if decision_clean == "KONFIRMASI_COCOK" else TIDAK_LULUS
        try:
            ev_query = (
                supabase.table("sync_evidence")
                .update({"evidence_status": evidence_new_status})
                .eq("result_id", result_id)
                .eq("evidence_status", PERLU_VERIFIKASI)
            )
            if target_clean in ("MOS", "MCF"):
                ev_query = ev_query.eq("category", target_clean)
            ev_query.execute()
        except Exception:
            pass

    # Recalculate session totals
    all_res = (
        supabase.table("sync_results")
        .select("validation_status, needs_review")
        .eq("sync_session_id", session_id)
        .execute()
    )
    items = all_res.data or []
    tot_valid = sum(1 for x in items if x.get("validation_status") == VALID)
    tot_nf = sum(1 for x in items if x.get("validation_status") == DATA_TIDAK_DITEMUKAN)
    tot_rev = sum(1 for x in items if x.get("needs_review") or x.get("validation_status") in (PERLU_VERIFIKASI, PENDING))

    supabase.table("sync_sessions").update({
        "total_valid": tot_valid,
        "total_review": tot_rev,
        "total_not_found": tot_nf,
    }).eq("id", session_id).execute()

    return {
        "result_id": result_id,
        "session_id": session_id,
        "decision": decision_clean,
        "note": note_clean,
        "previous_status": previous_status,
        "new_status": new_status,
        "mos_status": new_mos_status,
        "mcf_status": new_mcf_status,
        "needs_review": needs_review,
        "audit_id": audit_id,
        "created_at": now_iso,
    }


def generate_sync_export_workbook(
    supabase: Client,
    session_id: str,
    mode: str = "final",
) -> BytesIO:
    """Hasilkan workbook Wisuda dengan kolom sertifikasi terisi.

    Aturan yang disetujui 2026-10-23:
    - Hanya status VALID (LULUS) yang mengisi sel.
    - Format: "MOS (Lulus)" / "MCF (Lulus)" / "MTCNA (Lulus)".
    - Title: "MOS", "MCF", "MTCNA", atau kombinasinya ("MOS, MTCNA", "MOS, MCF", dst.).
    - Sel yang tidak lulus dibiarkan kosong.
    Mode 'draft' tetap tersedia sebagai pratinjau berlabel.
    """
    mode_clean = mode.strip().lower() if mode else "final"

    if mode_clean == "final" and not ADR_OUTPUT_MAPPING_APPROVED:
        raise ExportBlockedByPolicyError(
            "Ekspor final diblokir: menunggu pengesahan keputusan bisnis ITCC."
        )

    session = get_sync_session(supabase, session_id)

    res_resp = (
        supabase.table("sync_results")
        .select("*")
        .eq("sync_session_id", session_id)
        .order("source_row")
        .execute()
    )
    results = res_resp.data or []

    is_draft = mode_clean == "draft"
    prefix = "[DRAFT] " if is_draft else ""

    wb = Workbook()
    ws_wisuda = wb.active
    ws_wisuda.title = "Wisuda (DRAFT)" if is_draft else "Wisuda"

    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bold_font = Font(bold=True)

    headers = [
        "No",
        "Kode Prodi",
        "Program Studi",
        "NIM",
        "Nama",
        "Sertifikasi MOS",
        "Sertifikasi MCF/MTCNA",
        "Title",
    ]
    ws_wisuda.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws_wisuda.cell(row=1, column=col_num)
        cell.font = bold_font
        cell.fill = header_fill

    for idx, r in enumerate(results, start=1):
        fill = build_wisuda_fill_values(r.get("mos_status"), r.get("mcf_status"), r.get("mcf_programs"))
        nim_val = r.get("nim", "")
        kode_prodi_val = r.get("kode_prodi") or resolve_kode_prodi(nim_val)
        prodi_val = r.get("program_studi") or resolve_program_studi(kode_prodi_val)

        ws_wisuda.append(
            [
                idx,
                kode_prodi_val,
                prodi_val,
                nim_val,
                r.get("nama_wisuda", ""),
                f"{prefix}{fill['mos']}" if fill["mos"] else "",
                f"{prefix}{fill['mcf']}" if fill["mcf"] else "",
                f"{prefix}{fill['title']}" if fill["title"] else "",
            ]
        )

    for col_index, width in enumerate([6, 10, 28, 14, 34, 20, 20, 18], start=1):
        ws_wisuda.column_dimensions[chr(64 + col_index)].width = width
    ws_wisuda.freeze_panes = "A2"

    ws_audit = wb.create_sheet(title="Ringkasan")
    ws_audit.append(["PARAMETER", "NILAI"])
    ws_audit.cell(row=1, column=1).font = bold_font
    ws_audit.cell(row=1, column=2).font = bold_font

    audit_meta = [
        ("Status Berkas", "DRAFT - PRATINJAU" if is_draft else "FINAL"),
        ("Session ID", session.get("id")),
        ("Nama File Asli", session.get("file_name")),
        ("Batch SITASI Terpakai", session.get("sitasi_import_id")),
        ("Batch Certiport Terpakai", session.get("certiport_import_id")),
        ("Total Mahasiswa", session.get("total_rows")),
        ("Total Lulus", session.get("total_valid")),
        ("Total Tidak Lulus/Perlu Verifikasi", session.get("total_review")),
        ("Total Tidak Ditemukan", session.get("total_not_found")),
        ("Waktu Sinkronisasi", session.get("created_at")),
        ("Aturan Isi", "Hanya status LULUS yang mengisi kolom; sel lain kosong."),
    ]
    for param, val in audit_meta:
        ws_audit.append([param, str(val)])
    ws_audit.column_dimensions["A"].width = 34
    ws_audit.column_dimensions["B"].width = 60

    buf = BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    return buf
